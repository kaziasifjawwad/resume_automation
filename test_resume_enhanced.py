from docx import Document

# Create a new Document
doc = Document()

# Add content with variations and spelling differences
doc.add_heading('John Smith', 0)
doc.add_heading('Senior Software Engineer', level=1)

doc.add_heading('EXPERIENCE:', level=2)
doc.add_paragraph('• 6+ years of experience in Python3 development')
doc.add_paragraph('• Strong background in ReactJS and modern frontend frameworks')
doc.add_paragraph('• Experience with NodeJS backend development')
doc.add_paragraph('• Proficient in PostgreSQL database design and optimization')
doc.add_paragraph('• AWS cloud services experience (EC2, S3, RDS)')
doc.add_paragraph('• Docker containerization and Kubernetes orchestration')
doc.add_paragraph('• Machine Learning projects using TensorFlow and scikit-learn')
doc.add_paragraph('• Version control with Git and GitHub')
doc.add_paragraph('• CI/CD pipeline implementation with Jenkins')
doc.add_paragraph('• Agile development methodologies and Scrum')

doc.add_heading('SKILLS:', level=2)
doc.add_paragraph('• Python, JavaScript, TypeScript, SQL')
doc.add_paragraph('• React, Vue.js, Angular (basic)')
doc.add_paragraph('• Node.js, Express.js, FastAPI')
doc.add_paragraph('• PostgreSQL, MySQL, MongoDB')
doc.add_paragraph('• Docker, K8s, Jenkins')
doc.add_paragraph('• AWS, Azure, GCP')
doc.add_paragraph('• Git, GitHub, GitLab')
doc.add_paragraph('• Machine Learning, AI, Data Science')
doc.add_paragraph('• REST APIs, GraphQL, JSON')
doc.add_paragraph('• Linux, Bash scripting, PowerShell')

doc.add_heading('EDUCATION:', level=2)
doc.add_paragraph('• Master\'s in Computer Science')
doc.add_paragraph('• AWS Solutions Architect Certification')
doc.add_paragraph('• Docker Certified Associate')

# Save the document
doc.save('test_resume_enhanced.docx')
print("Enhanced test resume created: test_resume_enhanced.docx")

